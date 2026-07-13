"""
XLSX writer para a tese Piso Salarial Docente (reflexo de quinquênio + 6ª parte).

Layout reproduz a planilha modelo Pinheiro "QUINQUÊNIO E SEXTA PARTE"
(pasta docs/referencias/QQ PISO):
- Título (azul) + REQUERENTE (dourado), cabeçalho na linha 3, dados da 4 em diante.
- Colunas: Data | Piso | Total Vantagens | Qtde Quinq. | % | Dif. Quinq. |
  Tem 6ª? | Dif. 6ª Parte | Total Devido.
- Cada ano fecha com "13 SALARIO"; anos ATIVOS também com "1/3 de férias"
  (aposentado não tem terço de férias).
- Última linha: TOTAL BRUTO.
- Planilha protegida: só as colunas de entrada (Piso, Qtde Quinq., Tem 6ª?)
  ficam editáveis; fórmulas e cabeçalhos travados.

Fórmulas idênticas ao modelo:
    Total Vantagens = Piso
    %               = Qtde Quinq. * 5%
    Dif. Quinq.     = Total Vantagens * %
    Dif. 6ª Parte   = IF(Tem 6ª="Sim", (Total Vantagens + Dif. Quinq.)/6, 0)
    Total Devido    = Dif. Quinq. + Dif. 6ª Parte
"""

from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side, Protection
from openpyxl.comments import Comment
from openpyxl.worksheet.properties import PageSetupProperties

from src.teses.base_tese import BaseTese


AZUL = "0D1525"
DOURADO = "B38642"
CINZA = "F2F2F2"
AMARELO = "FFF2CC"
BRANCO = "FFFFFF"

FONT_NAME = "Lato"
MONEY_FMT = '_-"R$"\\ * #,##0.00_-;\\-"R$"\\ * #,##0.00_-;_-"R$"\\ * "-"??_-;_-@_-'
PCT_FMT = '0%'
DATE_FMT = 'mmm-yy'

thin = Border(
    left=Side(style='thin', color="D0D0D0"),
    right=Side(style='thin', color="D0D0D0"),
    top=Side(style='thin', color="D0D0D0"),
    bottom=Side(style='thin', color="D0D0D0"),
)

# Colunas (1-indexed)
COL_DATA = 1
COL_PISO = 2
COL_TOTAL_VANT = 3
COL_QTD_QUINQ = 4
COL_PCT = 5
COL_DIF_QUINQ = 6
COL_TEM_SEXTA = 7
COL_DIF_SEXTA = 8
COL_TOTAL_DEV = 9
LAST_COL = COL_TOTAL_DEV

HEADERS = {
    COL_DATA: "DATA DE\nPAGAMENTO",
    COL_PISO: "PISO SALARIAL\nDOCENTE",
    COL_TOTAL_VANT: "TOTAL VANTAGENS\nINTEGRAIS",
    COL_QTD_QUINQ: "QTDE.\nQUINQUÊNIOS",
    COL_PCT: "PORCENTAGEM",
    COL_DIF_QUINQ: "DIFERENÇA\nQUINQUÊNIOS",
    COL_TEM_SEXTA: "TEM 6ª\nPARTE?",
    COL_DIF_SEXTA: "DIFERENÇA\n6ª PARTE",
    COL_TOTAL_DEV: "TOTAL DEVIDO",
}

# Colunas de entrada (editáveis quando a planilha está protegida)
COLS_EDITAVEIS = {COL_PISO, COL_QTD_QUINQ, COL_TEM_SEXTA}
# Colunas de cálculo temporal (fundo cinza, como no modelo)
COLS_CINZA = {COL_TOTAL_VANT, COL_PCT, COL_DIF_QUINQ, COL_DIF_SEXTA}

amarelo_fill = PatternFill(start_color=AMARELO, end_color=AMARELO, fill_type="solid")
cinza_fill = PatternFill(start_color=CINZA, end_color=CINZA, fill_type="solid")
azul_fill = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")
dourado_fill = PatternFill(start_color=DOURADO, end_color=DOURADO, fill_type="solid")
unlocked = Protection(locked=False)


def write_piso_xlsx(resultado: dict, output_path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cálculo"

    nome_cliente = resultado['nome_cliente']
    periodos = resultado.get('periodos', {})
    situacao_por_periodo = resultado.get('situacao_por_periodo', {})

    _cabecalho(ws, nome_cliente)
    _corpo(ws, periodos, situacao_por_periodo)

    # Proteção: trava tudo, exceto as células de entrada (marcadas no corpo).
    # formatColumns=False libera o ajuste de largura das colunas mesmo com a
    # planilha protegida.
    ws.protection.sheet = True
    ws.protection.formatColumns = False

    _larguras(ws)
    _layout_impressao(ws)
    ws.freeze_panes = "A4"

    wb.save(output_path)
    return output_path


def _layout_impressao(ws):
    """Configura 'salvar como PDF' / impressão: paisagem, todas as colunas em
    uma página (largura), altura livre em várias páginas."""
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def _cabecalho(ws, nome_cliente: str):
    # Linha 1: título
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=LAST_COL)
    c = ws.cell(row=1, column=1,
                value="RECÁLCULO DE ADICIONAIS TEMPORAIS (QUINQUÊNIO E SEXTA "
                      "PARTE) SOBRE O PISO SALARIAL DOCENTE")
    c.font = Font(name=FONT_NAME, size=12, bold=True, color=BRANCO)
    c.fill = azul_fill
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30

    # Linha 2: requerente
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=LAST_COL)
    c = ws.cell(row=2, column=1, value=f"REQUERENTE: {nome_cliente}")
    c.font = Font(name=FONT_NAME, size=12, bold=True, color=BRANCO)
    c.fill = dourado_fill
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[2].height = 21

    # Linha 3: cabeçalhos
    for col, texto in HEADERS.items():
        c = ws.cell(row=3, column=col, value=texto)
        c.font = Font(name=FONT_NAME, size=9, bold=True, color=BRANCO)
        c.fill = azul_fill
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = thin
    ws.row_dimensions[3].height = 40


def _corpo(ws, periodos: dict, situacao_por_periodo: dict):
    sorted_p = sorted(periodos.keys())
    meses = _all_months(sorted_p)
    if not meses:
        return

    anos_ativos = _classificar_anos(meses, periodos, situacao_por_periodo)

    row = 4
    primeira_linha_dados = row
    ano_atual = None
    linha_inicio_ano = row
    ultima_linha = row

    for mes in meses:
        ano = int(mes[:4])

        if ano_atual is not None and ano != ano_atual:
            row = _fechar_ano(ws, row, linha_inicio_ano, ultima_linha,
                              anos_ativos.get(ano_atual, True))
            linha_inicio_ano = row
            ano_atual = ano
        elif ano_atual is None:
            ano_atual = ano
            linha_inicio_ano = row

        _linha_mes(ws, row, mes, periodos.get(mes))
        ultima_linha = row
        row += 1

    # Fecha o último ano
    row = _fechar_ano(ws, row, linha_inicio_ano, ultima_linha,
                      anos_ativos.get(ano_atual, True))

    _total_bruto(ws, row, primeira_linha_dados, row - 1)


def _linha_mes(ws, row, mes, dados):
    y, m = int(mes[:4]), int(mes[5:7])
    cell_a = ws.cell(row=row, column=COL_DATA, value=datetime(y, m, 1))
    cell_a.number_format = DATE_FMT
    cell_a.font = Font(name=FONT_NAME, size=10, bold=True)
    cell_a.alignment = Alignment(horizontal='center', vertical='center')
    cell_a.border = thin

    if not dados:
        # Mês sem holerite: só bordas (mantém o bloco anual coerente)
        for col in range(2, LAST_COL + 1):
            c = ws.cell(row=row, column=col)
            c.border = thin
            if col in COLS_CINZA:
                c.fill = cinza_fill
            elif col == COL_TOTAL_DEV:
                c.fill = amarelo_fill
            if col in COLS_EDITAVEIS:
                c.protection = unlocked
        return

    # B: Piso — fórmula auditável se houver atrasados; célula editável
    _cell_piso(ws, row, dados['normal'], dados['atrasados'])

    # C: Total Vantagens Integrais = Piso
    _cell_formula(ws, row, COL_TOTAL_VANT, f"=B{row}", fill=cinza_fill)

    # D: Qtde Quinquênios (entrada editável)
    cell_d = ws.cell(row=row, column=COL_QTD_QUINQ, value=dados['quinquenios'])
    cell_d.font = Font(name=FONT_NAME, size=10)
    cell_d.alignment = Alignment(horizontal='center', vertical='center')
    cell_d.border = thin
    cell_d.protection = unlocked

    # E: Porcentagem = D*5%
    _cell_formula(ws, row, COL_PCT, f"=D{row}*5%", fill=cinza_fill, fmt=PCT_FMT)
    # F: Diferença Quinquênios = C*E
    _cell_formula(ws, row, COL_DIF_QUINQ, f"=C{row}*E{row}", fill=cinza_fill)

    # G: Tem 6ª Parte? (entrada editável)
    cell_g = ws.cell(row=row, column=COL_TEM_SEXTA,
                     value="Sim" if dados.get('tem_sexta_parte') else "Não")
    cell_g.font = Font(name=FONT_NAME, size=10)
    cell_g.alignment = Alignment(horizontal='center', vertical='center')
    cell_g.border = thin
    cell_g.protection = unlocked

    # H: Diferença 6ª Parte = IF(G="Sim",(C+F)/6,0)
    _cell_formula(ws, row, COL_DIF_SEXTA,
                  f'=IF(G{row}="Sim",(C{row}+F{row})/6,0)', fill=cinza_fill)
    # I: Total Devido = F+H
    _cell_formula(ws, row, COL_TOTAL_DEV, f"=F{row}+H{row}",
                  fill=amarelo_fill, bold=True)


def _cell_piso(ws, row, normal, atrasados):
    """Coluna Piso (editável). Com atrasados: fórmula =normal+atr1+... amarela."""
    cell = ws.cell(row=row, column=COL_PISO)
    cell.number_format = MONEY_FMT
    cell.font = Font(name=FONT_NAME, size=10)
    cell.border = thin
    cell.alignment = Alignment(horizontal='right', vertical='center')
    cell.protection = unlocked

    if not atrasados:
        cell.value = normal if normal else None
        return cell

    parts = []
    if normal:
        parts.append(f"{normal:.2f}")
    for _, val in atrasados:
        parts.append(f"{val:.2f}")
    cell.value = "=" + ("+".join(parts) if parts else "0")
    cell.fill = amarelo_fill

    linhas = []
    if normal:
        linhas.append(f"Normal: R$ {normal:.2f}")
    for comp_pgto, val in atrasados:
        pgto = BaseTese.format_comp_display(BaseTese.mes_pagamento(comp_pgto))
        linhas.append(f"Atraso pago em {pgto}: R$ {val:.2f}")
    cell.comment = Comment("\n".join(linhas), "HoleritePRO")
    return cell


def _cell_formula(ws, row, col, formula, fill=None, fmt=None, bold=False):
    cell = ws.cell(row=row, column=col, value=formula)
    cell.number_format = fmt or MONEY_FMT
    cell.font = Font(name=FONT_NAME, size=10, bold=bold)
    cell.border = thin
    cell.alignment = Alignment(horizontal='right', vertical='center')
    if fill:
        cell.fill = fill
    return cell


def _fechar_ano(ws, row, linha_inicio, linha_fim, ano_ativo: bool) -> int:
    """13 SALARIO (sempre) + 1/3 de férias (só anos ativos)."""
    label_font = Font(name=FONT_NAME, size=10, bold=True)

    # 13 SALARIO
    _linha_fechamento(ws, row, "13 SALARIO", label_font,
                      f'=SUM(I{linha_inicio}:I{linha_fim})/12')
    linha_13 = row
    row += 1

    # 1/3 de férias — só ativo
    if ano_ativo:
        _linha_fechamento(ws, row, "1/3 de férias", label_font,
                          f'=I{linha_13}/3')
        row += 1

    return row


def _linha_fechamento(ws, row, label, label_font, formula_total):
    cell_a = ws.cell(row=row, column=COL_DATA, value=label)
    cell_a.font = label_font
    cell_a.alignment = Alignment(horizontal='center', vertical='center')
    cell_a.fill = amarelo_fill
    cell_a.border = thin
    for col in range(2, LAST_COL):
        c = ws.cell(row=row, column=col)
        c.border = thin
        c.fill = amarelo_fill
    cell_i = ws.cell(row=row, column=COL_TOTAL_DEV, value=formula_total)
    cell_i.number_format = MONEY_FMT
    cell_i.font = label_font
    cell_i.fill = amarelo_fill
    cell_i.border = thin
    cell_i.alignment = Alignment(horizontal='right', vertical='center')


def _total_bruto(ws, row, primeira_linha, ultima_linha):
    total_font = Font(name=FONT_NAME, size=11, bold=True)
    ws.merge_cells(start_row=row, start_column=COL_DIF_QUINQ,
                   end_row=row, end_column=COL_DIF_SEXTA)
    cell_lbl = ws.cell(row=row, column=COL_DIF_QUINQ, value="TOTAL BRUTO:")
    cell_lbl.font = total_font
    cell_lbl.fill = amarelo_fill
    cell_lbl.alignment = Alignment(horizontal='right', vertical='center')
    cell_lbl.border = thin

    cell_total = ws.cell(row=row, column=COL_TOTAL_DEV,
                         value=f'=SUM(I{primeira_linha}:I{ultima_linha})')
    cell_total.number_format = MONEY_FMT
    cell_total.font = total_font
    cell_total.fill = amarelo_fill
    cell_total.border = thin
    cell_total.alignment = Alignment(horizontal='right', vertical='center')


def _classificar_anos(meses, periodos, situacao_por_periodo) -> dict:
    """
    Classifica cada ano como ativo/inativo. Ano é ATIVO se teve ao menos um mês
    de holerite ativo (DDPE); INATIVO se só teve inativo (SPPREV). Anos sem
    informação herdam a classificação do ano anterior (transição ativo→inativo).
    Retorna {ano: True(ativo)/False(inativo)}.
    """
    anos = sorted({int(m[:4]) for m in meses})
    resultado = {}
    ultimo = True  # começa ativo por padrão
    for ano in anos:
        sits = [
            situacao_por_periodo.get(m)
            for m in meses
            if int(m[:4]) == ano and m in periodos
        ]
        if 'ativo' in sits:
            resultado[ano] = True
            ultimo = True
        elif 'inativo' in sits:
            resultado[ano] = False
            ultimo = False
        else:
            resultado[ano] = ultimo
    return resultado


def _all_months(sorted_periods: list) -> list:
    if not sorted_periods:
        return []
    y, m = int(sorted_periods[0][:4]), int(sorted_periods[0][5:7])
    ey, em = int(sorted_periods[-1][:4]), int(sorted_periods[-1][5:7])
    # Completa o último ano até dezembro se cair no ano corrente (linhas em
    # branco), mantendo o fechamento anual coerente.
    if ey == datetime.now().year:
        em = 12
    out = []
    while (y, m) <= (ey, em):
        out.append(f"{y:04d}-{m:02d}")
        m += 1
        if m > 12:
            m = 1
            y += 1
    return out


def _larguras(ws):
    larguras = {
        'A': 13, 'B': 18, 'C': 18, 'D': 13, 'E': 12,
        'F': 16, 'G': 12, 'H': 16, 'I': 16,
    }
    for col, w in larguras.items():
        ws.column_dimensions[col].width = w
