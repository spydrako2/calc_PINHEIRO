"""
Writer da planilha de IR sobre Bônus — reproduz a planilha modelo do
escritório, com as FÓRMULAS de cálculo e a aba "Dados" (tabelas IRPF) que elas
referenciam. O programa só lança os inputs; o Excel calcula e permite auditoria.

Uma aba "Cálculo" com dois blocos empilhados, como no modelo:
    BLOCO 1 — BONIFICAÇÃO POR RESULTADOS (RRA): colunas A..P, cálculo RRA.
    BLOCO 2 — ISENÇÃO (FUNDEB / DEJEC): colunas A..I, IR devido = 0
              (verba isenta) → restituição integral do IR pago.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.properties import PageSetupProperties

from src.core.irpf_rra_tables import TABELAS, DEDUCAO_POR_DEPENDENTE

AZUL = "0D1525"
DOURADO = "B38642"
CINZA = "F2F2F2"
AMARELO = "FFF2CC"
AZUL_CLARO = "EAF1FB"   # células de fórmula (não editar)
BRANCO = "FFFFFF"

FONT_NAME = "Lato"
# Zero é exibido como "R$ 0,00" (e não como traço), para não parecer célula vazia.
MONEY_FMT = '_-"R$"\\ * #,##0.00_-;\\-"R$"\\ * #,##0.00_-;_-"R$"\\ * #,##0.00_-;_-@_-'
PCT_FMT = "0.0%"
DATE_FMT = "dd/mm/yyyy"

thin = Border(
    left=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
    top=Side(style="thin", color="D0D0D0"),
    bottom=Side(style="thin", color="D0D0D0"),
)

ISENCAO_LABEL = {"FUNDEB": "ABONO FUNDEB", "DEJEC": "DEJEC"}

# Cabeçalho do bloco RRA (A..P)
HDR_RRA = [
    "PERÍODO\n(INÍCIO)", "PERÍODO\n(FIM)", "QTD.\nMESES", "DATA DE\nPAGAMENTO",
    "ANO-\nCALENDÁRIO", "BÔNUS PAGO\nEM FOLHA", "BASE RRA\nBRUTA", "Nº DE\nDEPENDENTES",
    "DEDUÇÃO POR\nDEPENDENTE", "BASE DE\nCÁLCULO RRA", "FAIXA SALARIAL\nTABELA IRPF-RRA",
    "ALÍQUOTA\nEFETIVA", "PARCELA\nDEDUZÍVEL", "IR DEVIDO", "IR PAGO\nEM FOLHA",
    "DIF.\nAPURADA",
]
# Cabeçalho do bloco Isenção (A..I)
HDR_ISENCAO = [
    "PERÍODO\n(INÍCIO)", "PERÍODO\n(FIM)", "QTD.\nMESES", "DATA DE\nPAGAMENTO",
    "ANO-\nCALENDÁRIO", "VALOR PAGO\nEM FOLHA", "IR DEVIDO", "IR PAGO\nEM FOLHA",
    "DIF.\nAPURADA",
]


def write_ir_bonus_xlsx(resultado: dict, output_path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cálculo"

    nome_cliente = resultado.get("nome_cliente", "")
    parcelas_rra = resultado.get("rra", [])
    isencao = resultado.get("isencao", {})

    # ---- Título geral (linha 1) e requerente (linha 2) ----
    _faixa(ws, 1, "RECÁLCULO DE IMPOSTO DE RENDA SOBRE BÔNUS", AZUL, 15, 28)
    _faixa(ws, 2, f"REQUERENTE: {nome_cliente}", DOURADO, 12, 20, align="left")

    row = 4
    dif_cells = []   # células de total de cada bloco (p/ VALOR DA CAUSA)

    if parcelas_rra:
        row, tot_cell = _bloco_rra(ws, row, parcelas_rra)
        dif_cells.append(tot_cell)
        row += 1

    for i, (tipo, parcelas) in enumerate(isencao.items(), start=2):
        if not parcelas:
            continue
        row, tot_cell = _bloco_isencao(ws, row, tipo, parcelas, ordem=i)
        dif_cells.append(tot_cell)
        row += 1

    if not parcelas_rra and not any(isencao.values()):
        ws["A4"] = "Nenhuma parcela de bônus encontrada nos holerites."

    # ---- VALOR DA CAUSA (soma dos totais dos blocos) ----
    if dif_cells:
        ws.merge_cells(f"A{row}:E{row}")
        c = ws[f"A{row}"]
        c.value = "VALOR DA CAUSA"
        c.font = Font(name=FONT_NAME, size=12, bold=True, color=BRANCO)
        c.fill = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")
        c.alignment = Alignment(horizontal="right", vertical="center")
        ct = ws[f"F{row}"]
        ct.value = "=" + "+".join(dif_cells)
        ct.number_format = MONEY_FMT
        ct.font = Font(name=FONT_NAME, size=12, bold=True, color=BRANCO)
        ct.fill = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")
        ct.alignment = Alignment(horizontal="right", vertical="center")

    _larguras(ws)
    _layout_impressao(ws)
    ws.freeze_panes = "A4"

    _dados_sheet(wb)
    wb.save(output_path)
    return output_path


def _bloco_rra(ws, row, parcelas):
    """Bloco 1: Bonificação por Resultados via RRA (colunas A..P)."""
    _titulo_bloco(ws, row, "PLANILHA DE CÁLCULO 1 — BONIFICAÇÃO POR RESULTADOS (RRA)", "P")
    row += 1
    _cabecalho(ws, row, HDR_RRA)
    row += 1

    amarelo = PatternFill(start_color=AMARELO, end_color=AMARELO, fill_type="solid")
    inicio = row
    for p in parcelas:
        _data(ws, row, 1, p["inicio"])
        _data(ws, row, 2, p["fim"])
        _formula(ws, row, 3, f'=IF(OR(A{row}="",B{row}=""),"",(YEAR(B{row})-YEAR(A{row}))*12+MONTH(B{row})-MONTH(A{row})+1)', fmt="0", center=True)
        _data(ws, row, 4, p["pagamento"])
        _formula(ws, row, 5, _ano_formula(row), fmt="General", center=True)
        _money(ws, row, 6, p["bonus"])
        _formula(ws, row, 7, f'=IF(OR(F{row}="",C{row}="",C{row}=0),"",F{row}/C{row})')
        _num(ws, row, 8, p.get("dependentes", 0))
        _formula(ws, row, 9, f'=IF(OR(H{row}="",H{row}=0),0,H{row}*Dados!$K$3)')
        _formula(ws, row, 10, f'=IF(G{row}="","",MAX(0,G{row}-I{row}))')
        _formula(ws, row, 11, _lookup(row, "P"), fmt="General")
        _formula(ws, row, 12, _lookup(row, "Q"), fmt=PCT_FMT)
        _formula(ws, row, 13, _lookup(row, "R"))
        _formula(ws, row, 14, f'=IF(OR(J{row}="",L{row}="",L{row}=0,C{row}=""),0,MAX(0,(J{row}*L{row}-M{row})*C{row}))')
        _money(ws, row, 15, p["ir_pago"])
        _formula(ws, row, 16, f'=IF(OR(O{row}="",N{row}=""),"",O{row}-N{row})', fill=amarelo, bold=True)
        row += 1

    tot_cell = _total_linha(ws, row, ultima=row - 1, primeira=inicio, col_dif=16, col_label_ate=15)
    return row, tot_cell


def _bloco_isencao(ws, row, tipo, parcelas, ordem):
    """Bloco de isenção (colunas A..I): IR devido = 0, restituição integral."""
    label = ISENCAO_LABEL.get(tipo, tipo)
    _titulo_bloco(ws, row, f"PLANILHA DE CÁLCULO {ordem} — {label} (ISENÇÃO — RESTITUIÇÃO INTEGRAL)", "I")
    row += 1
    _cabecalho(ws, row, HDR_ISENCAO)
    row += 1

    amarelo = PatternFill(start_color=AMARELO, end_color=AMARELO, fill_type="solid")
    inicio = row
    for p in parcelas:
        _data(ws, row, 1, p["inicio"])
        _data(ws, row, 2, p["fim"])
        _formula(ws, row, 3, f'=IF(OR(A{row}="",B{row}=""),"",(YEAR(B{row})-YEAR(A{row}))*12+MONTH(B{row})-MONTH(A{row})+1)', fmt="0", center=True)
        _data(ws, row, 4, p["pagamento"])
        _formula(ws, row, 5, _ano_formula(row), fmt="General", center=True)
        _money(ws, row, 6, p["bonus"])
        # G = IR devido = 0 (verba isenta)
        _num_money_zero(ws, row, 7)
        _money(ws, row, 8, p["ir_pago"])
        # I = diferença = IR pago − IR devido (= todo o IR pago)
        _formula(ws, row, 9, f'=IF(OR(H{row}="",G{row}=""),"",H{row}-G{row})', fill=amarelo, bold=True)
        row += 1

    tot_cell = _total_linha(ws, row, ultima=row - 1, primeira=inicio, col_dif=9, col_label_ate=8)
    return row, tot_cell


# ---------- helpers de célula ----------

def _faixa(ws, row, texto, cor, size, height, align="center"):
    ws.merge_cells(f"A{row}:P{row}")
    c = ws[f"A{row}"]
    c.value = texto
    c.font = Font(name=FONT_NAME, size=size, bold=True, color=BRANCO)
    c.fill = PatternFill(start_color=cor, end_color=cor, fill_type="solid")
    c.alignment = Alignment(horizontal=align, vertical="center")
    ws.row_dimensions[row].height = height


def _titulo_bloco(ws, row, texto, last_col):
    ws.merge_cells(f"A{row}:{last_col}{row}")
    c = ws[f"A{row}"]
    c.value = texto
    c.font = Font(name=FONT_NAME, size=11, bold=True, color=BRANCO)
    c.fill = PatternFill(start_color=DOURADO, end_color=DOURADO, fill_type="solid")
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 18


def _cabecalho(ws, row, titulos):
    fill = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")
    for i, titulo in enumerate(titulos, start=1):
        cell = ws.cell(row=row, column=i, value=titulo)
        cell.font = Font(name=FONT_NAME, size=9, bold=True, color=BRANCO)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
    ws.row_dimensions[row].height = 40


def _total_linha(ws, row, ultima, primeira, col_dif, col_label_ate):
    from openpyxl.utils import get_column_letter
    amarelo = PatternFill(start_color=AMARELO, end_color=AMARELO, fill_type="solid")
    last_label_col = get_column_letter(col_label_ate)
    ws.merge_cells(f"A{row}:{last_label_col}{row}")
    c = ws[f"A{row}"]
    c.value = "TOTAL A RECUPERAR"
    c.font = Font(name=FONT_NAME, size=11, bold=True)
    c.alignment = Alignment(horizontal="right", vertical="center")
    c.fill = amarelo
    c.border = thin
    dif_col = get_column_letter(col_dif)
    ct = ws.cell(row=row, column=col_dif)
    ct.value = f"=SUM({dif_col}{primeira}:{dif_col}{ultima})" if ultima >= primeira else 0
    ct.number_format = MONEY_FMT
    ct.font = Font(name=FONT_NAME, size=11, bold=True)
    ct.fill = amarelo
    ct.border = thin
    ct.alignment = Alignment(horizontal="right", vertical="center")
    return f"{dif_col}{row}"


def _data(ws, row, col, value):
    cell = ws.cell(row=row, column=col, value=value)
    cell.number_format = DATE_FMT
    cell.font = Font(name=FONT_NAME, size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin
    return cell


def _money(ws, row, col, value):
    cell = ws.cell(row=row, column=col, value=value)
    cell.number_format = MONEY_FMT
    cell.font = Font(name=FONT_NAME, size=10)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    cell.border = thin
    return cell


def _num(ws, row, col, value):
    cell = ws.cell(row=row, column=col, value=value)
    cell.number_format = "0"
    cell.font = Font(name=FONT_NAME, size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin
    return cell


def _num_money_zero(ws, row, col):
    # IR devido = 0 fixo (verba isenta) — tratado como fórmula/sistema (azul claro).
    cell = ws.cell(row=row, column=col, value=0)
    cell.number_format = MONEY_FMT
    cell.font = Font(name=FONT_NAME, size=10)
    cell.alignment = Alignment(horizontal="right", vertical="center")
    cell.border = thin
    cell.fill = PatternFill(start_color=AZUL_CLARO, end_color=AZUL_CLARO, fill_type="solid")
    return cell


def _formula(ws, row, col, formula, fmt=MONEY_FMT, fill=None, bold=False, center=False):
    cell = ws.cell(row=row, column=col, value=formula)
    cell.number_format = fmt
    cell.font = Font(name=FONT_NAME, size=10, bold=bold)
    cell.border = thin
    cell.alignment = Alignment(horizontal="center" if center else "right", vertical="center")
    # Fórmula → azul claro por padrão (sinaliza "não editar"); resultado usa amarelo.
    cell.fill = fill or PatternFill(start_color=AZUL_CLARO, end_color=AZUL_CLARO, fill_type="solid")
    return cell


def _ano_formula(r):
    return (f'=IF(D{r}="","",IF(D{r}<DATE(2023,5,1),"Ano_2015",IF(D{r}<DATE(2024,2,1),'
            f'"Ano_2023",IF(D{r}<DATE(2025,5,1),"Ano_2024","Ano_2025"))))')


def _lookup(r, col_dados):
    return (
        f'=IF(OR(J{r}="",E{r}=""),"",IFERROR(INDEX(Dados!${col_dados}$2:${col_dados}$21,'
        f'SUMPRODUCT(--(Dados!$M$2:$M$21=E{r})*--(Dados!$N$2:$N$21<=J{r})*'
        f'--(Dados!$O$2:$O$21>=J{r})*ROW(Dados!$M$2:$M$21))-1),""))'
    )


def _larguras(ws):
    widths = {"A": 12, "B": 12, "C": 7, "D": 13, "E": 12, "F": 13, "G": 12, "H": 11,
              "I": 13, "J": 13, "K": 22, "L": 10, "M": 12, "N": 13, "O": 13, "P": 14}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _dados_sheet(wb):
    """Aba Dados: tabelas IRPF-RRA no layout que as fórmulas referenciam
    (M2:R21) + dedução por dependente em K3."""
    ws = wb.create_sheet("Dados")
    hdr = Font(name=FONT_NAME, size=10, bold=True, color=BRANCO)
    hdr_fill = PatternFill(start_color=AZUL, end_color=AZUL, fill_type="solid")

    ws["K2"] = "DED. DEP"
    ws["K2"].font = hdr
    ws["K2"].fill = hdr_fill
    ws["K3"] = DEDUCAO_POR_DEPENDENTE
    ws["K3"].number_format = MONEY_FMT

    for col, titulo in (("M", "ANO_KEY"), ("N", "LIMITE_INF"), ("O", "LIMITE_SUP"),
                        ("P", "FAIXA_DESCR"), ("Q", "ALÍQUOTA"), ("R", "DEDUÇÃO")):
        ws[f"{col}1"] = titulo
        ws[f"{col}1"].font = hdr
        ws[f"{col}1"].fill = hdr_fill

    r = 2
    for ano in ("Ano_2015", "Ano_2023", "Ano_2024", "Ano_2025"):
        for faixa in TABELAS[ano]:
            ws[f"M{r}"] = ano
            ws[f"N{r}"] = faixa.limite_inf
            ws[f"O{r}"] = faixa.limite_sup
            ws[f"P{r}"] = faixa.descricao
            ws[f"Q{r}"] = faixa.aliquota
            ws[f"Q{r}"].number_format = PCT_FMT
            ws[f"R{r}"] = faixa.parcela_deduzir
            ws[f"R{r}"].number_format = MONEY_FMT
            r += 1

    for col in ("K", "M", "N", "O", "P", "Q", "R"):
        ws.column_dimensions[col].width = 16
    ws.column_dimensions["P"].width = 32
    return ws


def _layout_impressao(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.5
